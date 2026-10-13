"""Read reference files and export metadata comparisons."""

import csv
import io
import math
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .errors import DocumentError
from .schemas import CSV_COLUMNS, LEGACY_NOTICE_COLUMN

MAX_ROWS = 5000
MAX_BYTES = 20 * 1024 * 1024
ALIASES = {
    "agreement value": "Aggrement Value",
    "agreement start date": "Aggrement Start Date",
    "agreement end date": "Aggrement End Date",
    LEGACY_NOTICE_COLUMN.casefold(): "Termination Notice (Days)",
}


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, bool):
        raise DocumentError("Boolean cells are not valid metadata values.")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value)
    if len(text) > 32767:
        raise DocumentError("A value exceeds Excel's 32,767-character cell limit.")
    return text


def read_table(data: bytes, filename: str, *, allow_empty=False) -> list[dict]:
    """Read all labelled sheets, retaining original values and their source."""
    if len(data) > MAX_BYTES:
        raise DocumentError("Evaluation file exceeds 20 MB.")
    rows = []
    valid_header = False

    def consume(values, source, required=False):
        nonlocal valid_header
        iterator = iter(values)
        header = next(iterator, ())
        columns = [ALIASES.get(str(v).strip().casefold(), str(v).strip()) for v in header]
        canonical = {c.casefold(): c for c in CSV_COLUMNS}
        columns = [canonical.get(c.casefold(), c) for c in columns]
        if not set(CSV_COLUMNS).issubset(columns):
            if required:
                raise DocumentError(f"{source}: required columns are {', '.join(CSV_COLUMNS)}.")
            return
        if any(columns.count(c) != 1 for c in CSV_COLUMNS):
            raise DocumentError(f"{source}: duplicate required columns.")
        valid_header = True
        indices = {c: columns.index(c) for c in CSV_COLUMNS}
        for number, values in enumerate(iterator, 2):
            if number > MAX_ROWS + 1 or len(values) > 100:
                raise DocumentError(f"{source}: exceeds {MAX_ROWS} rows or 100 columns.")
            if not any(v is not None and str(v).strip() for v in values):
                continue
            if len(values) < max(indices.values()) + 1:
                raise DocumentError(f"{source}, row {number}: incomplete row.")
            if required and len(values) != len(columns):
                raise DocumentError(f"{source}, row {number}: malformed CSV row.")
            if any(isinstance(v, str) and v.startswith("=") for v in values):
                raise DocumentError(f"{source}, row {number}: replace formulas with saved values.")
            row = {c: cell_text(values[i]) for c, i in indices.items()}
            if not row["File Name"].strip():
                raise DocumentError(f"{source}, row {number}: missing filename.")
            row["ground_truth_source"] = f"{source}, row {number}"
            rows.append(row)
            if len(rows) > MAX_ROWS:
                raise DocumentError(f"Evaluation file exceeds {MAX_ROWS} rows.")

    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            consume(csv.reader(io.StringIO(data.decode("utf-8-sig"))), filename, True)
        elif suffix == ".xlsx":
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                    raise DocumentError("Expanded Excel file exceeds 100 MB.")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                for sheet in workbook:
                    if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                        raise DocumentError(f"Sheet {sheet.title} exceeds {MAX_ROWS} rows.")
                    if sheet.max_column and sheet.max_column > 100:
                        raise DocumentError(f"Sheet {sheet.title} exceeds 100 columns.")
                    consume(sheet.iter_rows(values_only=True), f"{filename} / {sheet.title}")
            finally:
                workbook.close()
        else:
            raise DocumentError("Evaluation supports .xlsx and .csv files.")
    except (UnicodeError, csv.Error, BadZipFile, ValueError, KeyError, OSError) as exc:
        raise DocumentError(f"Cannot read evaluation file: {filename}.") from exc
    if not rows and not (allow_empty and valid_header):
        raise DocumentError(f"{filename}: no data rows with the required columns.")
    return rows


def _comparison_result(detail: dict) -> tuple[str, str]:
    outcome = detail.get("comparison_status", "Mismatch")
    similarity = detail.get("fuzzy_score")
    score = f" ({similarity:.2f}% similarity)" if similarity is not None else ""
    if outcome.startswith("unlabelled"):
        return "Not evaluated", "No ground truth for this document; excluded from scores."
    if outcome == "Missing document":
        return "No", "No extraction result; incorrect in all three recall scores."
    if outcome == "Invalid value":
        return "No", "Extracted Metadata or ground truth has an invalid value."
    if outcome == "Review":
        reason = "Name requires review" + score
        if detail.get("fuzzy_conflicts"):
            reason += "; " + "; ".join(detail["fuzzy_conflicts"])
        return "No", reason + "; no automatic credit."
    if outcome == "Fuzzy match":
        return "Yes", "Fuzzy name match" + score + "; all identity safeguards passed."
    if detail.get("normalized_match"):
        if detail.get("exact_match"):
            return "Yes", "Exact match." if detail.get("prediction") else "Both values are empty."
        reason = (
            "Match after name cleaning."
            if detail["field"] in {"Party One", "Party Two"}
            else "Match after value normalization."
        )
        return "Yes", reason
    return "No", "Values differ after normalization." if detail.get(
        "prediction"
    ) else "Value not extracted."


def excel_report(report: dict, details: list[dict]) -> bytes:
    """Export computed scores and original strings; source text never becomes a formula."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Agreement extraction evaluation"])
    summary.append(["Metric", "Value"])
    for title, key in [
        ("Reference documents", "labelled_documents"),
        ("Scored documents (including failures)", "evaluated_documents"),
        ("Completed labelled documents", "completed_documents"),
        ("Unscored documents (no reference)", "unlabelled_documents"),
        ("Missing Extracted Metadata", "missing_documents"),
    ]:
        summary.append([title, report["overall"][key]])
    tiers = ("strict", "normalized", "tolerant")
    for tier in tiers:
        summary.append([f"Overall {tier} recall", report["overall"][tier]["recall"]])
    summary.append([])
    summary.append(
        [
            "Field",
            "Documents",
            "Strict correct",
            "Normalized correct",
            "Tolerant correct",
            "Strict recall",
            "Normalized recall",
            "Tolerant recall",
        ]
    )
    field_header = summary.max_row
    for name, score in report["fields"].items():
        summary.append(
            [
                name.replace("Aggrement", "Agreement"),
                score["total"],
                *[score[tier]["correct"] for tier in tiers],
                *[score[tier]["recall"] for tier in tiers],
            ]
        )
        for column in (6, 7, 8):
            summary.cell(summary.max_row, column).number_format = "0.0%"
    for row in summary.iter_rows(min_row=3, max_row=field_header - 2):
        if "recall" in str(row[0].value).casefold():
            row[1].number_format = "0.0%"
    summary.append([])
    notes_start = summary.max_row + 1
    for note in report["definitions"].values():
        summary.append([note])
    comparisons = workbook.create_sheet("Comparisons")
    comparisons.append(
        ["Document", "Field", "Extracted Metadata", "Ground truth",
         "Strict match", "Normalized match", "Tolerant match", "Comparison", "Reason"]
    )
    for detail in details:
        _, reason = _comparison_result(detail)
        matches = [
            "Not scored" if detail.get("excluded_from_metrics") else
            ("Yes" if detail.get(key) else "No")
            for key in ("exact_match", "normalized_match", "tolerant_match")
        ]
        comparisons.append(
            [
                detail["filename"],
                detail["field"].replace("Aggrement", "Agreement"),
                detail.get("prediction", ""),
                detail.get("ground_truth", ""),
                *matches,
                detail.get("comparison_status", "Mismatch"),
                reason,
            ]
        )
    for sheet in workbook:
        sheet.freeze_panes = "C2" if sheet == comparisons else f"B{field_header + 1}"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.print_area = sheet.dimensions
        if sheet != summary:
            sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    if len(cell.value) > 32767:
                        raise DocumentError("A value exceeds Excel's 32,767-character cell limit.")
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        headers = [1, 2, field_header] if sheet == summary else [1]
        for index in headers:
            for cell in sheet[index]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="17365D")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 26
        widths = (
            [38, 16, 17, 19, 17, 18, 20, 18] if sheet == summary
            else [34, 24, 32, 32, 16, 18, 18, 23, 55]
        )
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
        sheet.print_title_rows = f"1:{field_header}" if sheet == summary else "1:1"
        if sheet == comparisons:
            for row in sheet.iter_rows(min_row=2):
                for result in row[4:7]:
                    color = {"Yes": "E2EFDA", "No": "FCE4D6"}.get(result.value, "E7E6E6")
                    result.fill = PatternFill("solid", fgColor=color)
                    result.font = Font(bold=True, color="17365D")
        for cells in sheet:
            lines = max(
                math.ceil(
                    len(str(cell.value or ""))
                    / max(1, sheet.column_dimensions[cell.column_letter].width - 2)
                )
                for cell in cells
            )
            sheet.row_dimensions[cells[0].row].height = max(28, min(240, lines * 16))
    summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    summary.row_dimensions[1].height = 30
    for index in range(notes_start, summary.max_row + 1):
        summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=8)
        summary.row_dimensions[index].height = 40
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def evaluation_issue_message(issue: dict) -> str:
    """Explain an evaluation note with its affected document, field and source/value."""
    reasons = {
        "identical duplicate label": (
            "Duplicate reference row ignored; this document is evaluated once."
        ),
        "invalid ground truth": "The reference answer has an invalid value; check the reference.",
        "invalid prediction": "Extracted Metadata has an invalid value; check the extracted value.",
        "missing prediction": (
            "No Extracted Metadata was produced; this document counts as incorrect in all scores."
        ),
        "unlabelled prediction": "No reference answer was found; this document is not scored.",
        "missing ground truth": "No reference answer was found; this document is not scored.",
    }
    location = issue.get("filename", "Document")
    if issue.get("field"):
        location += " / " + issue["field"].replace("Aggrement", "Agreement")
    message = f"{location}: {reasons.get(issue['type'], 'Check the evaluation data.')}"
    if issue.get("details"):
        label = "Source" if issue["type"] == "identical duplicate label" else "Details"
        message += f" {label}: {issue['details']}"
    return message

