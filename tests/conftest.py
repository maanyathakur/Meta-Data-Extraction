import io
from dataclasses import replace

import pytest
from docx import Document
from PIL import Image

from metadata_extraction.config import Settings
from metadata_extraction.providers.base import ProviderResult
from metadata_extraction.schemas import FIELD_COLUMNS, AgreementMetadata


@pytest.fixture
def settings(tmp_path):
    return replace(Settings(), cache_dir=tmp_path / "cache", output_dir=tmp_path / "outputs")


@pytest.fixture
def docx_bytes():
    doc = Document()
    doc.add_paragraph("Agreement starts here.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Example Owner"
    doc.add_paragraph("Agreement ends here.")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def png_bytes():
    stream = io.BytesIO()
    Image.new("RGB", (100, 4000), "white").save(stream, format="PNG")
    return stream.getvalue()


@pytest.fixture
def missing_metadata():
    field = {
        "value": None,
        "status": "missing",
        "provenance": None,
        "evidence": [],
        "derivation": None,
        "explanation": "Not stated in this fixture.",
    }
    return AgreementMetadata.model_validate({name: field for name in FIELD_COLUMNS})


class FakeProvider:
    def __init__(self, metadata):
        self.metadata = metadata
        self.calls = 0

    def extract(self, document, system_prompt):
        self.calls += 1
        assert "Agreement" in system_prompt or "agreement" in system_prompt
        return ProviderResult(self.metadata, self.metadata.model_dump_json(), "fixture-model", {})


@pytest.fixture
def fake_provider(missing_metadata):
    return FakeProvider(missing_metadata)
