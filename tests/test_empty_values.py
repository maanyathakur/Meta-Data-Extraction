import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from metadata_extraction.api import create_app
from metadata_extraction.evaluation_io import excel_report
from metadata_extraction.evaluator import evaluate
from metadata_extraction.pipeline import ExtractionPipeline
from metadata_extraction.schemas import AgreementMetadata


@pytest.mark.parametrize("status", ["missing", "ambiguous"])
def test_unresolved_values_stay_empty_in_api_csv_and_excel(
    settings, fake_provider, docx_bytes, missing_metadata, status
):
    data = missing_metadata.model_dump()
    data["agreement_value"]["status"] = status
    metadata = AgreementMetadata.model_validate(data)
    fake_provider.metadata = metadata
    api = TestClient(create_app(settings, ExtractionPipeline(settings, fake_provider)))
    response = api.post("/extract?debug=true", files={"file": ("a.docx", docx_bytes)})
    assert response.status_code == 200
    result = response.json()
    assert result["values"]["agreement_value"] is None
    assert result["metadata"]["agreement_value"]["status"] == status
    assert result["metadata"]["agreement_value"]["explanation"]
    row = metadata.csv_row("a")
    assert row["Aggrement Value"] == ""
    report, details = evaluate([row], [row])
    workbook = load_workbook(io.BytesIO(excel_report(report, details)))
    try:
        sheet = workbook["Comparisons"]
        headings = [cell.value for cell in sheet[1]]
        assert sheet.cell(2, headings.index("Extracted Metadata") + 1).value in (None, "")
    finally:
        workbook.close()

