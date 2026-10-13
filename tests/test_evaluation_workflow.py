import base64
import io
import json
from dataclasses import replace
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from metadata_extraction.api import create_app
from metadata_extraction.dataset import read_labels
from metadata_extraction.errors import DocumentError
from metadata_extraction.evaluation_io import evaluation_issue_message, excel_report, read_table
from metadata_extraction.evaluator import csv_text, evaluate, evaluate_files
from metadata_extraction.schemas import CSV_COLUMNS, LEGACY_NOTICE_COLUMN


def row(name="a", **values):
    return {"File Name": name, **dict.fromkeys(CSV_COLUMNS[1:], ""), **values}


def xlsx(rows, title="Labels"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Notes"
    sheet.append(["Ignore this sheet"])
    sheet = workbook.create_sheet(title)
    sheet.append(CSV_COLUMNS)
    for item in rows:
        sheet.append([item[c] for c in CSV_COLUMNS])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_normalization_metrics_and_original_values():
    label = row(
        " A.PDF.DOCX ",
        **{
            "Aggrement Value": "12,000.00",
            "Aggrement Start Date": "01/04/2010",
            "Party One": " OWNER  NAME ",
            "Termination Notice (Days)": "30.0",
        },
    )
    prediction = row(
        "a.png",
        **{
            "Aggrement Value": "12000",
            "Aggrement Start Date": "2010-04-01",
            "Party One": "owner name",
            "Termination Notice (Days)": "30",
        },
    )
    report, details = evaluate([label], [prediction])
    assert report["overall"]["normalized"]["recall"] == 1
    assert report["overall"]["tolerant"]["recall"] == 1
    assert report["fields"]["Party One"]["normalized"]["recall"] == 1
    assert next(d for d in details if d["field"] == "Party One")["ground_truth"] == " OWNER  NAME "


def test_literal_invalid_values_are_flagged_and_missing_documents_count():
    label = row(**{"Aggrement End Date": "31.02.2011", "Party One": "Example Owner"})
    predicted = row(**{"Aggrement End Date": "31.02.2011", "Party One": "Example Owners"})
    report, details = evaluate([label, row("missing")], [predicted, row("extra")])
    assert report["fields"]["Aggrement End Date"]["strict"]["recall"] == 0.5
    assert any(issue["type"] == "invalid ground truth" for issue in report["data_issues"])
    assert report["fields"]["Party One"]["tolerant"]["recall"] == 0
    name = next(d for d in details if d["field"] == "Party One" and d["filename"] == "a")
    assert name["fuzzy_score"] > 90 and not name["normalized_match"]
    assert report["overall"]["missing_documents"] == 1
    assert report["overall"]["unlabelled_documents"] == 1
    assert len(details) == 18


def test_duplicates_merge_and_conflicts_fail():
    report, _ = evaluate([row("A.docx"), row("a.png")], [row()])
    assert report["label_rows"] == 1
    assert report["data_issues"][0]["type"] == "identical duplicate label"
    with pytest.raises(DocumentError, match="Conflicting ground truth"):
        evaluate([row(), row(**{"Party One": "different"})], [row()])
    with pytest.raises(DocumentError, match="Duplicate prediction"):
        evaluate([row()], [row("A.docx"), row("a.png")])


def test_large_numbers_do_not_round_or_expand_unbounded():
    report, _ = evaluate(
        [row(**{"Aggrement Value": "123456789012345678901234567890"})],
        [row(**{"Aggrement Value": "123456789012345678901234567891"})],
    )
    assert report["fields"]["Aggrement Value"]["normalized"]["recall"] == 0
    report, _ = evaluate([row(**{"Aggrement Value": "1e9999999"})], [row()])
    assert any(i["type"] == "invalid ground truth" for i in report["data_issues"])


def test_excel_input_dates_sources_and_safe_output():
    rows = read_table(
        xlsx(
            [
                row(
                    **{
                        "Aggrement Start Date": datetime(2020, 1, 2),
                        "Party One": "+untrusted",
                    }
                )
            ]
        ),
        "labels.xlsx",
    )
    assert rows[0]["Aggrement Start Date"] == "02.01.2020"
    assert "Labels, row 2" in rows[0]["ground_truth_source"]
    prediction = dict(rows[0], **{"Party One": "=malicious()"})
    report, details = evaluate(rows, [prediction])
    workbook = load_workbook(io.BytesIO(excel_report(report, details)))
    assert workbook.sheetnames == ["Summary", "Comparisons"]
    assert not any(
        "prediction" in str(cell.value).casefold()
        for sheet in workbook
        for cells in sheet
        for cell in cells
    )
    assert [c.value for c in workbook["Comparisons"][1]] == [
        "Document",
        "Field",
        "Extracted Metadata",
        "Ground truth",
        "Strict match", "Normalized match", "Tolerant match", "Comparison",
        "Reason",
    ]
    assert workbook["Summary"].max_column == 8
    summary_rows = list(workbook["Summary"].values)
    assert not any(
        "Every Yes in Comparisons counts as correct." in str(row) for row in summary_rows
    )
    assert (
        next(r for r in summary_rows if r[0] == "Overall strict recall")[1]
        == report["overall"]["strict"]["recall"]
    )
    assert not any(cell.data_type == "f" for sheet in workbook for cells in sheet for cell in cells)
    assert any(cell.value == "=malicious()" for cells in workbook["Comparisons"] for cell in cells)


@pytest.mark.parametrize(
    "content,name",
    [
        (b"not excel", "labels.xlsx"),
        (b"wrong,columns\na,b", "labels.csv"),
        (b"anything", "labels.xls"),
    ],
)
def test_bad_files_rejected(content, name):
    with pytest.raises(DocumentError):
        read_table(content, name)


def test_formulas_rejected():
    with pytest.raises(DocumentError, match="replace formulas"):
        read_table(xlsx([row(**{"Aggrement Value": "=100+200"})]), "labels.xlsx")


@pytest.mark.parametrize("legacy", [False, True])
def test_notice_column_compatibility_for_csv_and_excel(tmp_path, legacy):
    expected = row(**{"Termination Notice (Days)": "30"})
    header = LEGACY_NOTICE_COLUMN if legacy else "Termination Notice (Days)"
    content = csv_text([expected]).replace("Termination Notice (Days)", header).encode("utf-8")
    path = tmp_path / "labels.csv"
    path.write_bytes(content)
    assert read_labels(path)[0] == expected
    assert read_table(content, "labels.csv")[0]["Termination Notice (Days)"] == "30"
    workbook = load_workbook(io.BytesIO(xlsx([expected])))
    workbook["Labels"].cell(1, CSV_COLUMNS.index("Termination Notice (Days)") + 1).value = header
    stream = io.BytesIO()
    workbook.save(stream)
    labels = read_table(stream.getvalue(), "labels.xlsx")
    report, _ = evaluate(labels, [expected])
    assert report["fields"]["Termination Notice (Days)"]["strict"]["recall"] == 1


def test_multiple_files_cli_path(tmp_path):
    a, b, predictions = [tmp_path / f"{name}.csv" for name in ("a", "b", "predictions")]
    a.write_text(csv_text([row()]), encoding="utf-8")
    b.write_text(csv_text([row("b")]), encoding="utf-8")
    predictions.write_text(csv_text([row(), row("b")]), encoding="utf-8")
    report = evaluate_files([a, b], predictions, tmp_path / "report")
    assert report["overall"]["strict"]["recall"] == 1
    assert (tmp_path / "report/evaluation_report.xlsx").exists()


@pytest.mark.parametrize("run_count,missing", [(4, 0), (6, 2), (6, 6)])
def test_api_scores_all_labelled_run_documents_including_failures(settings, run_count, missing):
    client = TestClient(create_app(settings))
    ids = [f"doc{i}" for i in range(run_count)]
    predictions = [row(name) for name in ids[: run_count - missing]]
    response = client.post(
        "/evaluate",
        data={"document_ids": json.dumps(ids)},
        files=[
            ("predictions", ("predictions.csv", csv_text(predictions))),
            ("ground_truth", ("all_labels.csv", csv_text([row(f"doc{i}") for i in range(10)]))),
        ],
    )
    assert response.status_code == 200, response.text
    report = response.json()["report"]
    assert report["overall"]["labelled_documents"] == run_count
    assert report["overall"]["missing_documents"] == missing
    assert report["overall"]["evaluated_documents"] == run_count
    assert report["overall"]["completed_documents"] == run_count - missing
    book = load_workbook(io.BytesIO(base64.b64decode(response.json()["excel_base64"])))
    missing_rows = [r for r in book["Comparisons"].values if r[7] == "Missing document"]
    assert len(missing_rows) == missing * 6
    assert all(r[4:7] == ("No", "No", "No") for r in missing_rows)
    assert all("incorrect in all three" in r[8] for r in missing_rows)
    assert report["excluded_label_rows"] == 10 - run_count
    for score in report["fields"].values():
        assert score["total"] == run_count
        for tier in ("strict", "normalized", "tolerant"):
            assert score[tier]["recall"] == (run_count - missing) / run_count
            assert score[tier]["incorrect"] == missing


def test_scope_ignores_unrelated_conflicting_labels():
    report, details = evaluate(
        [row("selected"), row("other"), row("other", **{"Party One": "Different"})],
        [row("selected")],
        document_ids=["SELECTED.docx"],
    )
    assert report["label_rows"] == 1
    assert len(details) == 6
    assert report["requested_document_ids"] == ["selected"]


@pytest.mark.parametrize("scope", [[], {}, [""], [1]])
def test_invalid_scope_is_rejected(scope):
    with pytest.raises(DocumentError):
        evaluate([row("a")], [], document_ids=scope)


def test_api_evaluates_local_and_uploaded_files_without_provider(settings, tmp_path):
    data = tmp_path / "data"
    data.mkdir()
    (data / "a.xlsx").write_bytes(xlsx([row()]))
    client = TestClient(create_app(replace(settings, data_dir=data, output_dir=tmp_path / "out")))
    assert client.get("/evaluation-files").json()["files"] == ["a.xlsx"]
    response = client.post(
        "/evaluate",
        data={"local_files": '["a.xlsx"]'},
        files=[
            ("predictions", ("predictions.csv", csv_text([row(), row("b")]))),
            ("ground_truth", ("b.xlsx", xlsx([row("b")]))),
        ],
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["report"]["overall"]["strict"]["recall"] == 1
    content = base64.b64decode(result["excel_base64"])
    assert load_workbook(io.BytesIO(content)).sheetnames[0] == "Summary"
    assert (
        tmp_path / "out/evaluations" / result["report_id"] / "evaluation_report.xlsx"
    ).read_bytes() == content
    for selected in ('["../outside.xlsx"]', "{}", "[]"):
        response = client.post(
            "/evaluate",
            data={"local_files": selected},
            files={
                "predictions": ("predictions.csv", csv_text([row()])),
            },
        )
        assert response.status_code == 422


@pytest.mark.parametrize(
    "kind,reason",
    [
        ("identical duplicate label", "Duplicate reference row ignored"),
        ("invalid ground truth", "reference answer has an invalid value"),
        ("invalid prediction", "Extracted Metadata has an invalid value"),
        ("missing prediction", "counts as incorrect"),
        ("unlabelled prediction", "this document is not scored"),
        ("missing ground truth", "this document is not scored"),
    ],
)
def test_evaluation_notes_explain_the_document_field_and_reason(kind, reason):
    message = evaluation_issue_message(
        {
            "type": kind,
            "filename": "agreement",
            "field": "Aggrement Start Date",
            "details": "Reference workbook, row 12",
        }
    )
    assert "agreement / Agreement Start Date" in message
    assert reason in message and "Reference workbook, row 12" in message
    assert "prediction" not in message.casefold()

