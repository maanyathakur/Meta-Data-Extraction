import io
from dataclasses import replace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image

from metadata_extraction.document_loader import load_bytes
from metadata_extraction.errors import DocumentError


def test_preserves_paragraph_table_order(settings, docx_bytes):
    content = load_bytes(docx_bytes, "contract.docx", settings)
    text = [b.text for b in content.text_blocks]
    assert text == ["Agreement starts here.", "Owner | Example Owner", "Agreement ends here."]


def test_tall_image_sections_cover_full_height_with_overlap(settings, png_bytes):
    content = load_bytes(png_bytes, "contract.png", settings)
    assert [p.source_id for p in content.images] == [
        "image-1:y0-1800",
        "image-1:y1650-3450",
        "image-1:y3300-4000",
    ]
    assert [Image.open(io.BytesIO(p.data)).size for p in content.images] == [
        (100, 1800),
        (100, 1800),
        (100, 700),
    ]


def test_embedded_image_is_included(settings, png_bytes):
    doc = Document()
    doc.add_paragraph("Text alongside scan")
    doc.add_picture(io.BytesIO(png_bytes), width=Inches(1))
    stream = io.BytesIO()
    doc.save(stream)
    content = load_bytes(stream.getvalue(), "scan.docx", settings)
    assert content.images
    assert "embedded-" in content.images[0].source_id


@pytest.mark.parametrize(
    "data,name",
    [(b"", "empty.docx"), (b"bad", "bad.png"), (b"bad", "bad.docx"), (b"hello", "file.pdf")],
)
def test_rejects_invalid_input(settings, data, name):
    with pytest.raises(DocumentError):
        load_bytes(data, name, settings)


def test_image_limit_never_silently_truncates(settings, png_bytes):
    with pytest.raises(DocumentError, match="MAX_IMAGE_PARTS"):
        load_bytes(png_bytes, "scan.png", replace(settings, max_image_parts=1))


def _run(text):
    run, node = OxmlElement("w:r"), OxmlElement("w:t")
    node.text = text
    run.append(node)
    return run


@pytest.mark.parametrize("location", ["inline", "cell", "nested-cell"])
def test_content_controls_preserve_clause_text(settings, location):
    document = Document()
    control, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    control.append(content)
    if location == "inline":
        document.add_paragraph("Monthly rent is ")._p.append(control)
        content.append(_run("8000"))
    else:
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        if location == "nested-cell":
            cell = cell.add_table(rows=1, cols=1).cell(0, 0)
        paragraph = OxmlElement("w:p")
        paragraph.append(_run("Monthly rent is 8000"))
        content.append(paragraph)
        cell._tc.append(control)
    stream = io.BytesIO()
    document.save(stream)
    loaded = load_bytes(stream.getvalue(), "control.docx", settings)
    text = "\n".join(block.text for block in loaded.text_blocks)
    assert text.count("Monthly rent is 8000") == 1


def test_tracked_changes_use_final_view_for_text_and_images(settings, png_bytes):
    document = Document()
    paragraph = document.add_paragraph("Monthly rent is ")
    inserted, deleted = OxmlElement("w:ins"), OxmlElement("w:del")
    inserted.append(_run("8000"))
    deleted.append(_run("9000"))
    image_run = paragraph.add_run()
    image_run.add_picture(io.BytesIO(png_bytes))
    deleted.append(image_run._r)
    paragraph._p.extend([deleted, inserted])
    removed_paragraph = document.add_paragraph("Removed party")
    moved_from = OxmlElement("w:moveFrom")
    moved_from.append(removed_paragraph._p)
    document.element.body.append(moved_from)
    stream = io.BytesIO()
    document.save(stream)
    loaded = load_bytes(stream.getvalue(), "tracked.docx", settings)
    assert [block.text for block in loaded.text_blocks] == ["Monthly rent is 8000"]
    assert not loaded.images


def test_zero_overlap_can_be_configured_through_environment(monkeypatch, png_bytes):
    from metadata_extraction.config import Settings

    monkeypatch.setenv("IMAGE_TILE_OVERLAP", "0")
    config = Settings.from_env()
    loaded = load_bytes(png_bytes, "scan.png", config)
    assert config.tile_overlap == 0
    assert [part.source_id for part in loaded.images] == [
        "image-1:y0-1800", "image-1:y1800-3600", "image-1:y3600-4000",
    ]
