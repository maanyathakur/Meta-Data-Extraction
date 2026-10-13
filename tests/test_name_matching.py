import io

import pytest
from openpyxl import load_workbook

from metadata_extraction.evaluation_io import excel_report
from metadata_extraction.evaluator import evaluate
from metadata_extraction.name_matching import clean_name, name_outcome
from metadata_extraction.schemas import CSV_COLUMNS


@pytest.mark.parametrize(
    "left,right",
    [
        ("MR.K.Kuttan ", "k kuttan"),
        ("Shri V.K.Nataraj", "V K NATARAJ"),
        ("V.V.Ravi Kian,,", "V V Ravi Kian"),
        (".B.Kishore", "B Kishore"),
        ("Ingles, Jr.", "Ingles. Jr."),
        ("Mrs. Asha & Mr. Ramesh K.N.", "Asha and Ramesh K N"),
        ("Dr. O’Neil", "oneil"),
    ],
)
def test_formatting_cleanup(left, right):
    assert clean_name(left) == clean_name(right)


def test_meaningful_company_words_numbers_and_joint_separators_remain():
    assert clean_name("SRI VYSHNAVI DAIRY Private Ltd.") == "sri vyshnavi dairy private ltd"
    assert clean_name("SRI VYSHNAVI DAIRY") == "sri vyshnavi dairy"
    assert clean_name("Shri Industries Ltd.") == "shri industries ltd"
    assert clean_name("PROSPER 4 LESS, LLC") == "prosper 4 less llc"
    assert clean_name("A & B") != clean_name("A B")
    assert clean_name("A and/or B") == clean_name("A / B")
    assert clean_name("A or B") != clean_name("A and B")


@pytest.mark.parametrize(
    "expected,predicted",
    [
        ("GERALDINE Q. GALINATO", "GERALDINE O. GALINATO"),
        ("A.B. Sharma", "A.C. Sharma"),
        ("Geraldine O. Galinato", "Geraldine Galinato"),
    ],
)
def test_conflicting_or_missing_initials_never_automatically_match(expected, predicted):
    outcome, score, conflict = name_outcome(clean_name(expected), clean_name(predicted))
    assert conflict
    assert outcome == ("Review" if score > 90 else "Mismatch")


@pytest.mark.parametrize(
    "score,outcome",
    [
        (100, "Fuzzy match"),
        (97.00001, "Fuzzy match"),
        (97, "Review"),
        (90.00001, "Review"),
        (90, "Mismatch"),
    ],
)
def test_threshold_boundaries_use_unrounded_scores(monkeypatch, score, outcome):
    monkeypatch.setattr("metadata_extraction.name_matching.ratio", lambda *args: score)
    assert name_outcome("alice smith", "alice smyth")[0] == outcome


def test_recall_equals_accepted_comparison_matches_and_originals_unchanged():
    def row(name, value):
        return {"File Name": name, **dict.fromkeys(CSV_COLUMNS[1:], ""), "Party One": value}

    labels = [
        row("exact", "Mr. K.Kuttan"),
        row("fuzzy", "Christopher Alexander Johnson"),
        row("review", "GERALDINE Q. GALINATO"),
    ]
    predictions = [
        row("exact", "K Kuttan"),
        row("fuzzy", "Christopher Alexander Jonson"),
        row("review", "GERALDINE O. GALINATO"),
    ]
    report, details = evaluate(labels, predictions)
    score = report["fields"]["Party One"]
    assert score["strict"]["recall"] == 0
    assert score["normalized"]["recall"] == 1 / 3
    assert score["tolerant"]["recall"] == 2 / 3
    assert score["tolerant"]["correct"] == 2 and score["tolerant"]["incorrect"] == 1
    assert score["review_count"] == 1
    assert not {"accuracy", "tolerant_accuracy", "value_recall"}.intersection(score)
    names = [d for d in details if d["field"] == "Party One"]
    assert [d["comparison_status"] for d in names] == ["Normalized match", "Fuzzy match", "Review"]
    assert names[0]["ground_truth"] == "Mr. K.Kuttan"
    assert names[2]["match"] is False
    assert predictions[0]["Party One"] == "K Kuttan"
    book = load_workbook(io.BytesIO(excel_report(report, details)))
    headers = [c.value for c in book["Comparisons"][1]]
    assert headers == ["Document", "Field", "Extracted Metadata", "Ground truth",
                       "Strict match", "Normalized match", "Tolerant match", "Comparison", "Reason"]
    name_rows = [
        r for r in book["Comparisons"].iter_rows(min_row=2, values_only=True) if r[1] == "Party One"
    ]
    assert [r[4] for r in name_rows] == ["No", "No", "No"]
    assert [r[5] for r in name_rows] == ["Yes", "No", "No"]
    assert [r[6] for r in name_rows] == ["Yes", "Yes", "No"]
    assert sum(r[6] == "Yes" for r in name_rows) == score["tolerant"]["correct"]
    summary_row = next(r for r in book["Summary"].values if r[0] == "Party One")
    assert summary_row[1:5] == (3, 0, 1, 2)
    assert "name cleaning" in name_rows[0][8]
    assert "Fuzzy name match" in name_rows[1][8]
    assert "requires review" in name_rows[2][8] and "initials" in name_rows[2][8]


def test_conflicting_initials_above_97_still_require_review():
    outcome, score, conflict = name_outcome(
        clean_name("Alexandria Geraldine Catherine Q. Galinato"),
        clean_name("Alexandria Geraldine Catherine O. Galinato"),
    )
    assert score > 97 and conflict and outcome == "Review"


def test_fuzzy_rules_never_apply_to_numbers_dates_or_notice():
    label = {"File Name": "a", **dict.fromkeys(CSV_COLUMNS[1:], "")}
    label.update(
        {
            "Aggrement Value": "123456789",
            "Aggrement Start Date": "01.04.2020",
            "Termination Notice (Days)": "100",
        }
    )
    prediction = dict(
        label,
        **{
            "Aggrement Value": "123456788",
            "Aggrement Start Date": "02.04.2020",
            "Termination Notice (Days)": "101",
        },
    )
    report, details = evaluate([label], [prediction])
    for field in ("Aggrement Value", "Aggrement Start Date", "Termination Notice (Days)"):
        assert report["fields"][field]["tolerant"]["recall"] == 0
        detail = next(d for d in details if d["field"] == field)
        assert detail["comparison_status"] == "Mismatch" and detail["fuzzy_score"] is None

