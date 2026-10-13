import io
import json
from dataclasses import replace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from pydantic import ValidationError

from metadata_extraction.document_loader import load_bytes
from metadata_extraction.errors import DocumentError, ProviderError
from metadata_extraction.evaluation_io import read_table
from metadata_extraction.evaluator import csv_text, evaluate_files
from metadata_extraction.pipeline import ExtractionPipeline
from metadata_extraction.schemas import AgreementMetadata, FieldResult


def found(value, source="body-1", quote="Agreement starts here."):
    return FieldResult(
        value=value,
        status="found",
        provenance="explicit",
        evidence=[{"source_id": source, "quote": quote}],
        derivation=None,
        explanation=None,
    )


@pytest.mark.parametrize(
    "field,value",
    [
        ("agreement_start_date", "31.02.2025"),
        ("agreement_end_date", "29.02.2023"),
        ("agreement_start_date", "1.01.2025"),
        ("agreement_value", "-100"),
        ("agreement_value", "NaN"),
        ("agreement_value", "1e3"),
        ("agreement_value", "12,34"),
        ("termination_notice_days", "30.5"),
        ("termination_notice_days", "-1"),
    ],
)
def test_invalid_selected_values_rejected(missing_metadata, field, value):
    data = missing_metadata.model_dump()
    data[field] = found(value).model_dump()
    with pytest.raises(ValidationError):
        AgreementMetadata.model_validate(data)


@pytest.mark.parametrize(
    "field,value",
    [
        ("agreement_start_date", "29.02.2024"),
        ("agreement_value", "12,345.50"),
        ("agreement_value", "0"),
        ("termination_notice_days", "30"),
    ],
)
def test_valid_selected_values_accepted(missing_metadata, field, value):
    data = missing_metadata.model_dump()
    data[field] = found(value).model_dump()
    assert getattr(AgreementMetadata.model_validate(data), field).value == value


@pytest.mark.parametrize(
    "source,quote", [("invented", "Agreement starts here."), ("body-1", "Invented quotation")]
)
def test_unverifiable_evidence_is_not_cached_or_returned(
    settings, fake_provider, docx_bytes, source, quote
):
    fake_provider.metadata.party_one = found("Someone", source, quote)
    with pytest.raises(ProviderError, match="Evidence validation failed"):
        ExtractionPipeline(settings, fake_provider).extract_bytes(docx_bytes, "a.docx")
    assert not list(settings.cache_dir.glob("*.json"))


def test_normal_api_rejects_unverifiable_evidence(settings, fake_provider, docx_bytes):
    from fastapi.testclient import TestClient

    from metadata_extraction.api import create_app

    fake_provider.metadata.party_one = found("Someone", "invented")
    api = TestClient(create_app(settings, ExtractionPipeline(settings, fake_provider)))
    response = api.post("/extract", files={"file": ("a.docx", docx_bytes)})
    assert response.status_code == 502
    assert "Evidence validation failed" in response.json()["detail"]
    assert "Party One: source ID was not supplied" in response.json()["detail"]


def test_quote_failure_identifies_field_and_source_without_logging_content(
    settings, fake_provider, docx_bytes
):
    fake_provider.metadata.party_one = found("Someone", "body-1", "Private invented quote")
    with pytest.raises(ProviderError, match="Party One: quotation differs from source body-1"):
        ExtractionPipeline(settings, fake_provider).extract_bytes(docx_bytes, "a.docx")
    text = (settings.output_dir / "provider_events.jsonl").read_text(encoding="utf-8")
    event = json.loads(text.splitlines()[-1])
    assert event["event"] == "evidence_validation_failed"
    assert event["issues"] == [{
        "field": "party_one", "source_id": "body-1", "reason": "quote_mismatch"
    }]
    assert "Private invented quote" not in text and "Someone" not in text
    assert "Agreement starts here" not in text
    assert not list(settings.cache_dir.glob("*.json"))


def test_unknown_source_diagnostics_do_not_echo_model_text(
    settings, fake_provider, docx_bytes
):
    fake_provider.metadata.party_one = found("Someone", "private-model-text", "private-quote")
    with pytest.raises(ProviderError) as failure:
        ExtractionPipeline(settings, fake_provider).extract_bytes(docx_bytes, "a.docx")
    text = (settings.output_dir / "provider_events.jsonl").read_text(encoding="utf-8")
    assert "private-model-text" not in str(failure.value) + text
    assert "private-quote" not in str(failure.value) + text
    assert json.loads(text.splitlines()[-1])["issues"][0]["source_id"] is None


def test_unique_verbatim_quote_corrects_only_its_source_id(
    settings, fake_provider, docx_bytes
):
    fake_provider.metadata.party_one = found("Someone", "body-1", "Example Owner")
    record = ExtractionPipeline(settings, fake_provider).extract_bytes(docx_bytes, "a.docx")
    evidence = record.metadata.party_one.evidence[0]
    assert evidence.source_id == "body-2"
    assert evidence.quote == "Example Owner"
    assert record.metadata.party_one.value == "Someone"
    assert fake_provider.metadata.party_one.evidence[0].source_id == "body-1"
    events = (settings.output_dir / "provider_events.jsonl").read_text(encoding="utf-8")
    assert json.loads(events.splitlines()[-1])["event"] == "evidence_source_corrected"
    assert "Example Owner" not in events
    cached = ExtractionPipeline(settings, fake_provider).extract_bytes(docx_bytes, "a.docx")
    assert cached.cached and fake_provider.calls == 1


def test_repeated_quote_does_not_guess_a_source(settings, fake_provider):
    document = Document()
    document.add_paragraph("Unrelated clause")
    document.add_paragraph("Repeated source text")
    document.add_paragraph("Repeated source text")
    stream = io.BytesIO()
    document.save(stream)
    fake_provider.metadata.party_one = found("Someone", "body-1", "Repeated source text")
    with pytest.raises(ProviderError, match="quotation differs from source body-1"):
        ExtractionPipeline(settings, fake_provider).extract_bytes(stream.getvalue(), "a.docx")
    assert not list(settings.cache_dir.glob("*.json"))


def test_cache_grounding_is_rechecked(settings, fake_provider, docx_bytes):
    pipeline = ExtractionPipeline(settings, fake_provider)
    pipeline.extract_bytes(docx_bytes, "a.docx")
    path = next(settings.cache_dir.glob("*.json"))
    record = json.loads(path.read_text())
    record["metadata"]["party_one"] = found("Someone", "invented").model_dump()
    path.write_text(json.dumps(record), encoding="utf-8")
    result = pipeline.extract_bytes(docx_bytes, "a.docx")
    assert not result.cached and result.metadata.party_one.value is None
    assert fake_provider.calls == 2


def test_header_footer_images_and_textbox_are_ingested(settings, png_bytes):
    doc = Document()
    doc.add_paragraph("Body")
    for region in (doc.sections[0].header, doc.sections[0].footer):
        region.paragraphs[0].add_run().add_picture(io.BytesIO(png_bytes))
    # A legacy VML textbox: w:pict/v:shape/v:textbox/w:txbxContent/w:p.
    run = doc.add_paragraph().add_run()
    pict = OxmlElement("w:pict")
    from lxml import etree

    shape = etree.SubElement(pict, "{urn:schemas-microsoft-com:vml}shape")
    textbox = etree.SubElement(shape, "{urn:schemas-microsoft-com:vml}textbox")
    content = OxmlElement("w:txbxContent")
    paragraph = OxmlElement("w:p")
    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Monthly rent is 9000."
    text_run.append(text)
    paragraph.append(text_run)
    content.append(paragraph)
    textbox.append(content)
    run._r.append(pict)
    stream = io.BytesIO()
    doc.save(stream)
    loaded = load_bytes(stream.getvalue(), "regions.docx", settings)
    assert any("header" in p.source_id for p in loaded.images)
    assert any("footer" in p.source_id for p in loaded.images)
    assert sum(b.text == "Monthly rent is 9000." for b in loaded.text_blocks) == 1


def test_empty_results_count_labelled_failures_as_zero_recall(tmp_path):
    from metadata_extraction.schemas import CSV_COLUMNS

    labels = tmp_path / "labels.csv"
    predictions = tmp_path / "predictions.csv"
    empty = csv_text([]).encode()
    labels.write_text(csv_text([{"File Name": "a", **dict.fromkeys(CSV_COLUMNS[1:], "")}]))
    predictions.write_bytes(empty)
    report = evaluate_files(labels, predictions, tmp_path / "report")
    for tier in ("strict", "normalized", "tolerant"):
        assert report["overall"][tier]["recall"] == 0
    assert report["overall"]["evaluated_documents"] == 1
    assert report["missing_prediction_rows"] == ["a"]
    with pytest.raises(DocumentError):
        read_table(empty, "labels.csv")
    with pytest.raises(DocumentError):
        read_table(b"invalid\n", "predictions.csv", allow_empty=True)


def test_combined_run_preserves_summary_when_all_requests_fail(
    tmp_path, settings, docx_bytes, missing_metadata, monkeypatch
):
    import importlib.util

    from metadata_extraction.config import PROJECT_ROOT
    from metadata_extraction.errors import ProviderQuotaError

    spec = importlib.util.spec_from_file_location(
        "combined_run", PROJECT_ROOT / "scripts/run_and_evaluate.py"
    )
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    config = replace(settings, api_key="synthetic", data_dir=tmp_path / "data")
    for split in ("train", "test"):
        folder = config.data_dir / split
        folder.mkdir(parents=True)
        (folder / "a.docx").write_bytes(docx_bytes)
        (config.data_dir / f"{split}.csv").write_text(csv_text([missing_metadata.csv_row("a")]))

    class Exhausted:
        def extract(self, *args):
            raise ProviderQuotaError(60)

    monkeypatch.setattr(runner.Settings, "from_env", lambda: config)
    monkeypatch.setattr("metadata_extraction.providers.create_provider", lambda _: Exhausted())
    assert runner.main() == 1
    summary = json.loads((config.output_dir / "run_summary.json").read_text())
    assert summary["finished_at"]
    assert all(s["run"]["failed_documents"] == 1 for s in summary["splits"].values())
